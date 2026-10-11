#!/usr/bin/env python3
"""Run quick structural checks against the real-layout catalogue."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

try:
    from pptx import Presentation
except ImportError as exc:  # pragma: no cover
    raise SystemExit("Missing dependency: python-pptx. Install it with: python3 -m pip install -r requirements.txt") from exc


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pptx", required=True, help="Generated PowerPoint")
    parser.add_argument("--catalog", required=True, help="layouts.json generated from the same template")
    args = parser.parse_args()

    catalog_data = json.loads(Path(args.catalog).expanduser().resolve().read_text(encoding="utf-8"))
    layouts = catalog_data.get("layouts", {})
    presentation = Presentation(str(Path(args.pptx).expanduser().resolve()))
    errors: list[str] = []
    warnings: list[str] = []
    slide_width, slide_height = presentation.slide_width, presentation.slide_height

    for slide_number, slide in enumerate(presentation.slides, start=1):
        layout_name = slide.slide_layout.name
        layout = layouts.get(layout_name)
        if layout is None:
            errors.append(f"Slide {slide_number}: layout {layout_name!r} is not in the catalogue")
            continue
        placeholders = layout.get("placeholders", {})
        for shape in slide.shapes:
            if not shape.is_placeholder:
                errors.append(f"Slide {slide_number}: non-placeholder shape {shape.name!r} was added")
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > slide_width or shape.top + shape.height > slide_height:
                errors.append(f"Slide {slide_number}: shape {shape.name!r} extends beyond the slide bounds")
            if not shape.is_placeholder:
                continue

            idx = str(shape.placeholder_format.idx)
            if idx == "11":
                continue  # Fixed copyright/page-number/footer placeholder in the brand master.
            metadata = placeholders.get(idx)
            if metadata is None:
                errors.append(f"Slide {slide_number}: placeholder idx={idx} is absent from the catalogue")
                continue
            if not shape.has_text_frame or not shape.text.strip():
                continue
            capacity = metadata.get("estimated_capacity_chars")
            if capacity and len(shape.text) > capacity:
                errors.append(
                    f"Slide {slide_number}, placeholder idx={idx}: {len(shape.text)} characters "
                    f"exceed estimated capacity {capacity}"
                )
            fill_ratio = len(shape.text) / capacity if capacity else None
            threshold = metadata.get("min_fill_ratio")
            if fill_ratio is not None and threshold is not None and fill_ratio < threshold:
                warnings.append(
                    f"Slide {slide_number}, placeholder idx={idx}: estimated fill "
                    f"{fill_ratio:.0%} is below {threshold:.0%}; review whether this layout is too dense"
                )
            floor = metadata.get("font_size_pt")
            if floor:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.font.size and run.font.size.pt < floor:
                            errors.append(
                                f"Slide {slide_number}, placeholder idx={idx}: "
                                f"font {run.font.size.pt:g}pt is below layout size {floor:g}pt"
                            )

    for warning in warnings:
        print(f"WARNING: {warning}")
    for error in errors:
        print(f"ERROR: {error}")
    if errors:
        print(f"Deck lint failed: {len(errors)} error(s), {len(warnings)} warning(s)")
        return 1
    print(f"Deck lint passed: {len(presentation.slides)} slides, {len(warnings)} layout-fit warning(s)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
