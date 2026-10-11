#!/usr/bin/env python3
"""Extract a fast, approximate placeholder catalogue from a PowerPoint template."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.enum.shapes import PP_PLACEHOLDER
except ImportError as exc:  # pragma: no cover
    raise SystemExit("Missing dependency: python-pptx. Install it with: python3 -m pip install -r requirements.txt") from exc


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", required=True, help="Source .pptx/.potx")
    parser.add_argument("--output", required=True, help="Path for layouts.json")
    args = parser.parse_args()

    presentation = Presentation(str(Path(args.template).expanduser().resolve()))
    result = {
        "schema_version": 1,
        "template": str(Path(args.template).name),
        "capacity_note": "Character capacity and 40% fill checks are heuristic warnings, not rendering guarantees.",
        "font": "Noto IKEA Latin",
        "colors": {
            "blue": "#0058A3",
            "yellow": "#FFDB00",
            "black": "#000000",
            "white": "#FFFFFF",
            "light_grey": "#F0F0F0",
            "grey": "#575757",
            "orange": "#F54C22",
            "green": "#36C336"
        },
        "layouts": {}
    }
    for layout in presentation.slide_masters[0].slide_layouts:
        content_placeholders = [
            shape for shape in layout.placeholders
            if str(shape.placeholder_format.idx) != "11"
            and (shape.has_text_frame or shape.placeholder_format.type == PP_PLACEHOLDER.PICTURE)
        ]
        placeholders = {}
        for shape in content_placeholders:
            kind = shape.placeholder_format.type
            is_image = kind == PP_PLACEHOLDER.PICTURE
            font_size = next(
                (
                    int(size) / 100
                    for size in shape._element.xpath(".//a:lvl1pPr/a:defRPr/@sz")
                ),
                None,
            ) if shape.has_text_frame else None
            estimate_font = font_size or (36 if str(kind) == "TITLE (1)" else 18)
            estimated_capacity = None
            if shape.has_text_frame:
                columns = max(1, int((shape.width / 914400 * 72) / (estimate_font * 0.52)))
                lines = max(1, int((shape.height / 914400 * 72) / (estimate_font * 1.25)))
                estimated_capacity = columns * lines
            placeholders[str(shape.placeholder_format.idx)] = {
                "name": shape.name,
                "type": str(kind),
                "x": shape.left,
                "y": shape.top,
                "width": shape.width,
                "height": shape.height,
                "text": shape.has_text_frame,
                "image": is_image,
                "font_size_pt": font_size or (36 if str(kind) == "TITLE (1)" else 18 if shape.has_text_frame else None),
                "estimated_capacity_chars": estimated_capacity,
                "min_fill_ratio": 0.4 if shape.has_text_frame else None,
            }
        result["layouts"][layout.name] = {"placeholders": placeholders}

    output = Path(args.output).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Wrote {output}: {len(result['layouts'])} real layouts")
    return 0


if __name__ == "__main__":
    sys.exit(main())
