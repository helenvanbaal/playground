#!/usr/bin/env python3
"""Generate a PowerPoint using a real layout and its native placeholders.

Usage:
    python3 generate_deck.py --outline outline.json --output outputs/deck.pptx
    python3 generate_deck.py --inspect-template
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Pt
except ImportError as exc:  # pragma: no cover
    raise SystemExit(
        "Missing dependency: python-pptx. Install it with: "
        "python3 -m pip install -r requirements.txt"
    ) from exc


DEFAULT_TEMPLATE_PATH = Path(__file__).resolve().parents[1] / "assets" / "base-template.pptx"
DEFAULT_CATALOG_PATH = Path(__file__).resolve().parents[1] / "layouts.json"
COPYRIGHT_TEXT = "© Ingka Holding B.V. 2026"
COPYRIGHT_COLOR = RGBColor(0x57, 0x57, 0x57)
COPYRIGHT_FONT_SIZE_PT = 6
BRAND_FONT = "Noto IKEA Latin"


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", help=f"PowerPoint template (default: {DEFAULT_TEMPLATE_PATH})")
    parser.add_argument("--catalog", help=f"Layout catalogue (default: {DEFAULT_CATALOG_PATH})")
    parser.add_argument("--outline", help="Path to a JSON outline")
    parser.add_argument("--output", help="Path to write the generated .pptx")
    parser.add_argument("--inspect-template", action="store_true", help="Print real layout names and placeholders")
    args = parser.parse_args()

    template_path = Path(args.template).expanduser().resolve() if args.template else DEFAULT_TEMPLATE_PATH
    if not template_path.exists():
        raise SystemExit(f"Template not found: {template_path}")
    presentation = Presentation(str(template_path))
    normalize_copyright(presentation)

    if args.inspect_template:
        inspect_template(presentation)
        return 0
    if not args.outline or not args.output:
        raise SystemExit("--outline and --output are required unless --inspect-template is used")

    catalog_path = Path(args.catalog).expanduser().resolve() if args.catalog else DEFAULT_CATALOG_PATH
    catalog = load_catalog(catalog_path)
    outline_path = Path(args.outline).expanduser().resolve()
    outline = load_json(outline_path, "outline")
    build_deck(presentation, outline, catalog)

    output_path = Path(args.output).expanduser().resolve()
    if output_path.suffix.lower() != ".pptx":
        raise SystemExit("Output path must end with .pptx")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(output_path))
    print(f"Wrote {output_path}")
    return 0


def normalize_copyright(presentation) -> None:
    """Give the inherited copyright line one fixed, subdued style in every deck."""
    matches = []
    for master in presentation.slide_masters:
        matches.extend(
            shape for shape in master.shapes
            if shape.has_text_frame and "© Ingka Holding B.V." in shape.text
        )
    if len(matches) != 1:
        raise SystemExit(f"Expected one master copyright text box; found {len(matches)}")

    copyright_shape = matches[0]
    text_frame = copyright_shape.text_frame
    text_frame.clear()
    paragraph = text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = COPYRIGHT_TEXT
    run.font.name = BRAND_FONT
    run.font.size = Pt(COPYRIGHT_FONT_SIZE_PT)
    run.font.color.rgb = COPYRIGHT_COLOR

    for master in presentation.slide_masters:
        for layout in master.slide_layouts:
            for shape in list(layout.shapes):
                if shape.has_text_frame and "© Ingka Holding B.V." in shape.text:
                    shape._element.getparent().remove(shape._element)
            if layout.name == "Title - Black":
                layout._element.set("showMasterSp", "1")


def load_json(path: Path, label: str) -> dict[str, Any]:
    if not path.is_file():
        raise SystemExit(f"{label.capitalize()} not found: {path}")
    try:
        with path.open(encoding="utf-8") as handle:
            data = json.load(handle)
    except json.JSONDecodeError as exc:
        raise SystemExit(f"Invalid {label} JSON at {path}: {exc}") from exc
    if not isinstance(data, dict):
        raise SystemExit(f"{label.capitalize()} must be a JSON object")
    return data


def load_catalog(path: Path) -> dict[str, Any]:
    catalog = load_json(path, "layout catalogue")
    layouts = catalog.get("layouts")
    if not isinstance(layouts, dict):
        raise SystemExit("Layout catalogue must contain a 'layouts' object")
    return layouts


def inspect_template(presentation) -> None:
    for layout in presentation.slide_masters[0].slide_layouts:
        print(f"=== {layout.name} ===")
        for shape in layout.placeholders:
            placeholder = shape.placeholder_format
            print(
                f"  idx={placeholder.idx} type={placeholder.type} "
                f"name={shape.name!r} x={shape.left} y={shape.top} "
                f"w={shape.width} h={shape.height}"
            )


def build_deck(presentation, outline: dict[str, Any], catalog: dict[str, Any]) -> None:
    slides = outline.get("slides")
    if not isinstance(slides, list) or not slides:
        raise SystemExit("Outline must contain a non-empty 'slides' array")

    # A source template may retain sample slides for reference. They are not
    # output slides; generated slides use the same real layouts and placeholders.
    slide_ids = presentation.slides._sldIdLst
    for slide_id in list(slide_ids):
        presentation.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)

    layouts = {layout.name: layout for layout in presentation.slide_masters[0].slide_layouts}
    for number, slide_data in enumerate(slides, start=1):
        if not isinstance(slide_data, dict):
            raise SystemExit(f"Slide {number} must be an object")
        layout_name = slide_data.get("layout")
        if not isinstance(layout_name, str) or layout_name not in layouts:
            raise SystemExit(
                f"Slide {number}: layout must be an exact real template layout name; "
                f"valid names: {', '.join(layouts)}"
            )
        if layout_name not in catalog:
            raise SystemExit(f"Slide {number}: layout {layout_name!r} is missing from {DEFAULT_CATALOG_PATH.name}")
        slide = presentation.slides.add_slide(layouts[layout_name])
        populate_slide(slide, slide_data, catalog[layout_name], number)


def populate_slide(slide, slide_data: dict[str, Any], layout_info: dict[str, Any], number: int) -> None:
    placeholders = {str(shape.placeholder_format.idx): shape for shape in slide.placeholders}
    text_values = slide_data.get("text", {})
    image_values = slide_data.get("images", {})
    if not isinstance(text_values, dict) or not isinstance(image_values, dict):
        raise SystemExit(f"Slide {number}: 'text' and 'images' must be objects keyed by placeholder index")

    known_placeholders = {str(key): value for key, value in layout_info.get("placeholders", {}).items()}
    for idx, text in text_values.items():
        key = str(idx)
        shape = placeholders.get(key)
        metadata = known_placeholders.get(key)
        if shape is None or metadata is None:
            raise SystemExit(f"Slide {number}: placeholder idx={key} does not exist on this layout")
        if not metadata.get("text"):
            raise SystemExit(f"Slide {number}: placeholder idx={key} is not a text placeholder")
        if not isinstance(text, str):
            raise SystemExit(f"Slide {number}: text for placeholder idx={key} must be a string")
        capacity = metadata.get("estimated_capacity_chars")
        if capacity and len(text) > capacity:
            raise SystemExit(
                f"Slide {number}: placeholder idx={key} is over estimated capacity "
                f"({len(text)} > {capacity} characters)"
            )
        shape.text = text

    for idx, image in image_values.items():
        key = str(idx)
        shape = placeholders.get(key)
        metadata = known_placeholders.get(key)
        if shape is None or metadata is None or not metadata.get("image"):
            raise SystemExit(f"Slide {number}: placeholder idx={key} is not an image placeholder")
        image_path = image.get("path") if isinstance(image, dict) else image
        if not isinstance(image_path, str):
            raise SystemExit(f"Slide {number}: image for placeholder idx={key} must be a path or object with path")
        resolved_path = Path(image_path).expanduser().resolve()
        if not resolved_path.is_file():
            raise SystemExit(f"Slide {number}: image not found: {resolved_path}")
        shape.insert_picture(str(resolved_path))

    notes = slide_data.get("notes")
    if notes:
        if not isinstance(notes, str):
            raise SystemExit(f"Slide {number}: notes must be a string")
        slide.notes_slide.notes_text_frame.text = notes

    unexpected = set(slide_data) - {"layout", "key_message", "text", "images", "notes"}
    if unexpected:
        raise SystemExit(
            f"Slide {number}: unsupported fields {sorted(unexpected)}; "
            "put content in the real layout's 'text' and 'images' placeholders"
        )


if __name__ == "__main__":
    sys.exit(main())
